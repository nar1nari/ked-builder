from docx.document import Document as DocumentObject
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx import Document

from ked_builder.utils import find_table_after_heading

HEADING_RU = "Модель выпускника и квалификационная"
HEADING_KK = "Түлек моделі және біліктілік сипаттамасы"

FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(11)

# В таблице "Модель выпускника ..." (и в ОП, и в шаблоне) 4 раздела идут в
# фиксированном порядке: сфера, объекты, предметы, виды профессиональной
# деятельности. В исходной таблице ОП это одна колонка, где строки
# чередуются: строка-заголовок раздела (индексы 0,2,4,6), затем строка с
# текстом (индексы 1,3,5,7) - отсюда и шаг 2 при переборе.
_FIELDS = ("sphere", "objects", "subjects", "types")


def get_professional_activity(docx_path: str, heading_text: str) -> dict[str, str]:
    """
    Возвращает {"sphere": ..., "objects": ..., "subjects": ..., "types": ...} -
    текст четырёх разделов ("Сфера/Объекты/Предметы/Виды профессиональной
    деятельности") из таблицы, идущей сразу после heading_text в документе ОП.
    """
    doc = Document(docx_path)

    table = find_table_after_heading(doc, heading_text)
    if table is None:
        raise ValueError("Таблица после заголовка не найдена")

    result: dict[str, str] = {}
    for i, field in enumerate(_FIELDS):
        row_index = 1 + i * 2
        if row_index < len(table.rows):
            result[field] = table.rows[row_index].cells[0].text.strip()

    return result


def _apply_font(run) -> None:
    run.bold = False
    run.font.name = FONT_NAME
    run.font.size = FONT_SIZE
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attr), FONT_NAME)


def _set_cell_text(cell, text: str) -> None:
    """
    Заполняет ячейку текстом, где каждая строка исходного текста (в ОП они
    разделены переносом строки - фактически это пункты списка) становится
    отдельным абзацем-пунктом списка, начинающимся с длинного тире "—".
    """
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    lines = [line.strip() for line in text.split("\n") if line.strip()]
    if not lines:
        lines = [""]

    paragraphs = cell.paragraphs
    # первый абзац ячейки уже существует - переиспользуем его для первой строки,
    # остальные строки получают новые абзацы через add_paragraph()
    first_p = paragraphs[0]
    for run in list(first_p.runs):
        run._element.getparent().remove(run._element)
    for extra_p in paragraphs[1:]:
        extra_p._p.getparent().remove(extra_p._p)

    for i, line in enumerate(lines):
        p = first_p if i == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        prefix = "— " if line else ""
        run = p.add_run(f"{prefix}{line}")
        _apply_font(run)


def fill_professional_activity_table(
    doc: DocumentObject, ru_docx_path: str, kk_docx_path: str
) -> None:
    """
    Находит в doc (шаблоне) таблицу "Кәсіби қызмет саласы / Сфера
    профессиональной деятельности" и заполняет её 4 пустые строки текстом
    из русского и казахского документов ОП: первый столбец - казахский
    текст, второй - русский (сама таблица ищется по тексту первой ячейки,
    а не по индексу - на случай, если порядок таблиц в шаблоне изменится).
    """
    ru_data = get_professional_activity(ru_docx_path, HEADING_RU)
    kk_data = get_professional_activity(kk_docx_path, HEADING_KK)

    target_table = None
    for table in doc.tables:
        if table.rows and "Кәсіби қызмет саласы" in table.rows[0].cells[0].text:
            target_table = table
            break
    if target_table is None:
        raise ValueError('Таблица "Кәсіби қызмет саласы" не найдена в шаблоне')

    for i, field in enumerate(_FIELDS):
        row_index = 1 + i * 2
        if row_index >= len(target_table.rows):
            continue
        row = target_table.rows[row_index]
        _set_cell_text(row.cells[0], kk_data.get(field, ""))
        _set_cell_text(row.cells[1], ru_data.get(field, ""))
