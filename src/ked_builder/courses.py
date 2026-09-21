import re
import difflib

import openpyxl
from docx import Document

from ked_builder.utils import find_table_after_heading


class Course:
    def __init__(
        self,
        name: str | None = None,
        name_kk: str | None = None,
        code: str | None = None,
        code_kk: str | None = None,
        credits: int | float | None = None,
        cycle: str | None = None,
        cycle_kk: str | None = None,
        module: str | None = None,
        module_kk: str | None = None,
        number: str | None = None,
        prerequisites: list[str] | None = None,
        postrequisites: list[str] | None = None,
        goal: str | None = None,
        goal_kk: str | None = None,
        summary: str | None = None,
        summary_kk: str | None = None,
        results: list[str] | None = None,
        mentor: str | None = None,
        department: str | None = None,
        component: str | None = None,
    ):
        self.name = name
        self.name_kk = name_kk
        self.code = code
        self.code_kk = code_kk
        self.credits = credits
        self.cycle = cycle
        self.cycle_kk = cycle_kk
        self.module = module
        self.module_kk = module_kk
        self.number = number
        self.prerequisites = prerequisites
        self.postrequisites = postrequisites
        self.goal = goal
        self.goal_kk = goal_kk
        self.summary = summary
        self.summary_kk = summary_kk
        self.results = results
        self.mentor = mentor
        self.department = department
        # тип компонента дисциплины, например "Обязательный компонент",
        # "Вузовский компонент" или "Компонент по выбору" - берётся из
        # строки-разделителя раздела в таблице "Сведения о дисциплинах
        # образовательной программы"
        self.component = component

    def __repr__(self):
        return (
            f"Course(name={self.name!r}, name_kk={self.name_kk!r}, code={self.code!r}, "
            f"code_kk={self.code_kk!r}, credits={self.credits!r}, number={self.number!r}, "
            f"component={self.component!r}, prerequisites={self.prerequisites!r}, "
            f"postrequisites={self.postrequisites!r})"
        )

    @property
    def is_elective(self) -> bool:
        """Является ли дисциплина компонентом по выбору (элективом)."""
        return bool(self.component) and "выбор" in self.component.lower()


HEADING_COURSES = "Сведения о дисциплинах образовательной программы"
HEADING_COURSES_KK = "Білім беру бағдарламасының пәндері туралы мәліметтер"
HEADING_PREREQUISITES = "Карта соответствия пререквизитов и постреквизитов дисциплин"


def _normalize_name(name: str) -> str:
    """
    Приводит название дисциплины к упрощённому виду для сравнения:
    нижний регистр, "," -> ".", убираем приставку "Minor 1/2 ..." и
    лишние пробелы. Нужно, потому что в разных документах одно и то же
    название дисциплины иногда записано по-разному (например
    "Психология, Культурология" и "Психология. Культурология", или
    "Minor 1 Организация бизнеса" и просто "Организация бизнеса").
    """
    n = name.lower().strip()
    n = n.replace(",", ".")
    n = re.sub(r"^minor\s*\d+\.?\s*", "", n)
    n = re.sub(r"\s+", " ", n)
    return n.strip(" .")


def _lookup(name: str, exact: dict, normalized: dict, normalized_keys: list[str]):
    """
    Ищет значение для названия дисциплины в справочнике, полученном из
    другого документа. Порядок попыток:
    1. точное совпадение названия,
    2. совпадение после нормализации (регистр, "," vs ".", "Minor ..."),
    3. ближайшее по написанию название (спасает от опечаток/небольших
       расхождений формулировок между учебным планом и документами ОП).
    """
    if name in exact:
        return exact[name]

    norm = _normalize_name(name)
    if norm in normalized:
        return normalized[norm]

    match = difflib.get_close_matches(norm, normalized_keys, n=1, cutoff=0.65)
    if match:
        return normalized[match[0]]

    return None


# =====================================================================
# Учебный план (УП, .xlsx) - основной источник истины: название, код,
# количество кредитов, цикл и модуль дисциплины на русском и казахском
# (в файле также есть английский, но он не используется).
# =====================================================================

# ячейка кода дисциплины (последняя, английская часть) заканчивается
# числом, опционально с "(номер семестровой части)" в скобках -
# например "1101" или "1103 (1)"
_TRAILING_NUMBER = re.compile(r"(\d[\d\s()]*)\s*$")


def _split_trilingual(text: str) -> tuple[str, str, str]:
    """
    Делит ячейку учебного плана вида "KZ-текст/ RU-текст/ EN-текст" на
    3 части. В реальном файле встречаются лишние пробелы и переносы
    строк вокруг "/", поэтому каждая часть дополнительно нормализуется.
    """
    parts = [re.sub(r"\s+", " ", p).strip() for p in text.split("/")]
    while len(parts) < 3:
        parts.append("")
    return parts[0], parts[1], parts[2]


def _split_code(text: str) -> tuple[str, str]:
    """
    Код дисциплины в учебном плане записан как
    "KZ-приставка/ RU-приставка/ EN-приставка НОМЕР" - например
    "KТ/ IK/ HK 1101". Номер один на все три языка, но физически
    дописан только к последней (английской) части. Возвращает
    (код на русском, код на казахском).
    """
    kk_prefix, ru_prefix, en_part = _split_trilingual(text)
    match = _TRAILING_NUMBER.search(en_part)
    number = re.sub(r"\s+", "", match.group(1)) if match else ""
    ru_code = f"{ru_prefix} {number}".strip()
    kk_code = f"{kk_prefix} {number}".strip()
    return ru_code, kk_code


# строка - дисциплина, если код дисциплины заканчивается на число
# (опционально с "(N)"); у строк-заголовков разделов ("1 курс",
# "Первый семестр", итоговых строк и т.п.) кода нет вообще, а у строки
# с заголовками столбцов код есть, но это не число - оба случая так
# отсеиваются
_CODE_ROW_PATTERN = re.compile(r"\d\s*(\(\d+\))?\s*$")


def get_curriculum_courses(xlsx_path: str) -> list[Course]:
    """
    Читает учебный план (.xlsx) и возвращает список дисциплин с уже
    заполненными name/name_kk/code/code_kk/credits/cycle/cycle_kk/
    module/module_kk. Остальные поля (goal, summary, results,
    prerequisites, component, ...) заполняются позже в get_courses() по
    данным из документов ОП (.docx).

    Модуль (столбец B) в файле объединяется по нескольким дисциплинам
    подряд - в openpyxl это выглядит как None у всех строк, кроме первой
    в группе, поэтому используется дотягивание последнего известного
    значения (last_module_raw).

    Количество кредитов для альтернатив внутри одного "Компонента по
    выбору" в файле не дублируется, а помечено "*" ("столько же, сколько
    у первой дисциплины этой группы выбора") - в этом случае берётся
    последнее известное числовое значение кредитов.
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active

    courses: list[Course] = []
    last_module_raw: str | None = None
    last_credits = None

    for row in ws.iter_rows(min_row=7):
        cycle_raw = row[0].value
        module_raw = row[1].value
        code_raw = row[2].value
        name_raw = row[3].value
        credits_raw = row[4].value

        if module_raw:
            last_module_raw = module_raw

        if not code_raw or not name_raw or not _CODE_ROW_PATTERN.search(str(code_raw)):
            continue  # заголовок раздела/итоговая строка - не дисциплина

        if credits_raw in (None, "*"):
            credits = last_credits
        else:
            credits = credits_raw
            last_credits = credits_raw

        cycle_kk, cycle_ru, _ = _split_trilingual(cycle_raw or "")
        module_kk, module_ru, _ = _split_trilingual(last_module_raw or "")
        name_kk, name_ru, _ = _split_trilingual(name_raw)
        code_ru, code_kk = _split_code(code_raw)

        courses.append(
            Course(
                name=name_ru,
                name_kk=name_kk,
                code=code_ru,
                code_kk=code_kk,
                credits=credits,
                cycle=cycle_ru,
                cycle_kk=cycle_kk,
                module=module_ru,
                module_kk=module_kk,
            )
        )

    return courses


# =====================================================================
# Документы образовательной программы (.docx) - источник для цели/
# краткого содержания, результатов обучения, типа компонента и карты
# пре-/постреквизитов. Сопоставляются с дисциплинами из учебного плана
# НЕ по точному совпадению названия (см. _lookup).
# =====================================================================


def _split_goal_summary(text: str) -> tuple[str | None, str | None]:
    """
    Делит полный текст описания дисциплины на:
      - goal: первое предложение (всё до первой точки включительно),
      - summary: всё, что осталось после неё.
    """
    text = text.strip()
    if not text:
        return None, None

    dot = text.find(".")
    if dot == -1:
        return text, None

    goal = text[: dot + 1].strip()
    summary = text[dot + 1 :].strip() or None
    return goal, summary


def get_course_details(
    docx_path: str, heading_text: str, with_results_and_component: bool = False
) -> dict[str, dict]:
    """
    Из таблицы "Сведения о дисциплинах ..." (на любом языке, определяется
    heading_text) собирает для каждой дисциплины {"goal": ..., "summary": ...},
    а если with_results_and_component=True - ещё "results" (список кодов
    результатов обучения, например ["ON1","ON2"]) и "component" (тип
    компонента дисциплины). Результаты обучения и тип компонента одинаковы
    независимо от языка описания, поэтому их достаточно взять один раз
    из русского документа.
    """
    doc = Document(docx_path)

    table = find_table_after_heading(doc, heading_text)
    if table is None:
        raise ValueError("Таблица после заголовка не найдена")

    details: dict[str, dict] = {}
    current_component = None

    for row in table.rows:
        cells = row.cells
        if len(cells) < 4:
            continue

        row_number = cells[0].text.strip()
        if not row_number.isdigit():
            # строка-разделитель раздела: например
            # "Цикл базовых дисциплин\nКомпонент по выбору" -
            # компонент - это последняя строка текста
            current_component = cells[0].text.strip().split("\n")[-1]
            continue

        name = cells[2].text.strip()
        goal, summary = _split_goal_summary(cells[3].text.strip())

        if name and name not in details:
            entry = {"goal": goal, "summary": summary}
            if with_results_and_component:
                entry["results"] = cells[1].text.strip().split()
                entry["component"] = current_component
            details[name] = entry

    return details


# строка-номер дисциплины вида "1-1", "4-3", либо просто "1" (для
# "Школьного курса" из нулевого семестра, у которого номер в таблице
# не содержит дефиса)
_NUMBER_PATTERN = re.compile(r"^\d+(-\d+)?$")


def get_course_numbers(docx_path: str) -> dict[str, dict]:
    """
    Из таблицы "Карта соответствия пререквизитов и постреквизитов дисциплин"
    собирает для каждой дисциплины:
        - number: её номер в карте (например "1-1")
        - prerequisites: список номеров дисциплин-пререквизитов
        - postrequisites: список номеров дисциплин-постреквизитов

    Возвращает словарь {название дисциплины: {number, prerequisites, postrequisites}}.
    """
    doc = Document(docx_path)

    table = find_table_after_heading(doc, HEADING_PREREQUISITES)
    if table is None:
        raise ValueError("Таблица после заголовка не найдена")

    entries: dict[str, dict] = {}
    for row in table.rows:
        cells = row.cells
        if len(cells) != 10:
            continue  # необычная строка (например строка про диплом) - пропускаем

        raw_number = cells[2].text.strip()
        if not _NUMBER_PATTERN.match(raw_number):
            continue  # это строка-разделитель семестра, а не дисциплина

        semester = cells[0].text.strip()
        number = raw_number if "-" in raw_number else f"{semester}-{raw_number}"

        name = cells[3].text.strip()
        prereq_raw = cells[6].text.strip()
        postreq_raw = cells[8].text.strip()

        def parse_list(raw: str) -> list[str]:
            if raw in ("", "-"):
                return []
            return [item.strip() for item in re.split(r",\s*", raw) if item.strip()]

        if name and name not in entries:
            entries[name] = {
                "number": number,
                "prerequisites": parse_list(prereq_raw),
                "postrequisites": parse_list(postreq_raw),
            }

    return entries


def get_courses(
    xlsx_path: str,
    ru_docx_path: str,
    kk_docx_path: str | None = None,
) -> list[Course]:
    """
    Собирает итоговый список дисциплин:
    - name/name_kk/code/code_kk/credits/cycle/cycle_kk/module/module_kk -
      из учебного плана (xlsx_path), это основной источник истины;
    - goal/summary/results/component/number/prerequisites/postrequisites -
      из русского документа ОП (ru_docx_path);
    - goal_kk/summary_kk - из казахского документа ОП (kk_docx_path), если передан.

    Между документами дисциплины сопоставляются НЕ по точному совпадению
    названия (в учебном плане и документах ОП формулировки могут немного
    отличаться), а через _lookup (точное -> нормализованное -> ближайшее
    по написанию совпадение).
    """
    courses = get_curriculum_courses(xlsx_path)

    ru_details = get_course_details(ru_docx_path, HEADING_COURSES, with_results_and_component=True)
    normalized_ru_details = {_normalize_name(name): entry for name, entry in ru_details.items()}
    normalized_ru_keys = list(normalized_ru_details.keys())

    numbers = get_course_numbers(ru_docx_path)
    normalized_numbers = {_normalize_name(name): entry for name, entry in numbers.items()}
    normalized_number_keys = list(normalized_numbers.keys())

    kk_details: dict[str, dict] = {}
    normalized_kk_details: dict[str, dict] = {}
    normalized_kk_keys: list[str] = []
    if kk_docx_path:
        kk_details = get_course_details(kk_docx_path, HEADING_COURSES_KK)
        normalized_kk_details = {_normalize_name(name): entry for name, entry in kk_details.items()}
        normalized_kk_keys = list(normalized_kk_details.keys())

    for course in courses:
        ru_entry = _lookup(course.name, ru_details, normalized_ru_details, normalized_ru_keys)
        if ru_entry:
            course.goal = ru_entry["goal"]
            course.summary = ru_entry["summary"]
            course.results = ru_entry["results"]
            course.component = ru_entry["component"]

        number_entry = _lookup(course.name, numbers, normalized_numbers, normalized_number_keys)
        if number_entry:
            course.number = number_entry["number"]
            course.prerequisites = number_entry["prerequisites"]
            course.postrequisites = number_entry["postrequisites"]

        if kk_docx_path and course.name_kk:
            kk_entry = _lookup(course.name_kk, kk_details, normalized_kk_details, normalized_kk_keys)
            if kk_entry:
                course.goal_kk = kk_entry["goal"]
                course.summary_kk = kk_entry["summary"]

    return courses
