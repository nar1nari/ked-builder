import difflib
import re
from os import PathLike
from typing import BinaryIO

import openpyxl

from ked_builder.utils import normalize_text

ExcelSource = str | PathLike | BinaryIO

# Порог схожести (0..1) для нечёткого поиска. Для преподавателей он
# намеренно строже, чем при сопоставлении описаний дисциплин: ошибочно
# подставленная фамилия хуже, чем отсутствующая (её видно в документе
# как красное "None" и можно заполнить вручную).
PROGRAM_CUTOFF = 0.8
DISCIPLINE_CUTOFF = 0.85


class FacultyIndex:
    """
    Преподаватели одной образовательной программы из файла ППС:
    {название дисциплины -> список ФИО}. Названия дисциплин в файле могут
    быть как на русском, так и на казахском (по языку группы), поэтому
    поиск по русскому названию находит преподавателей русскоязычных
    групп, а по казахскому - казахскоязычных.
    """

    def __init__(self, teachers_by_discipline: dict[str, list[str]]):
        # ключи - нормализованные названия (см. utils.normalize_text)
        self._teachers = teachers_by_discipline
        self._keys = list(teachers_by_discipline.keys())

    def find(self, discipline_name: str | None) -> list[str]:
        """
        Возвращает преподавателей дисциплины или [] (если не найдено).
        Название ищется не по точному совпадению: сначала по нормализованному
        тексту (регистр, пробелы, знаки препинания не важны), затем по
        ближайшему по написанию названию (спасает от опечаток в файле).
        """
        if not discipline_name:
            return []

        norm = normalize_text(discipline_name)
        if not norm:
            return []
        if norm in self._teachers:
            return list(self._teachers[norm])

        match = difflib.get_close_matches(
            norm, self._keys, n=1, cutoff=DISCIPLINE_CUTOFF
        )
        return list(self._teachers[match[0]]) if match else []

    def find_mentor(self, discipline_name: str | None) -> str | None:
        """Как find(), но одной строкой через запятую; None - если никого не нашли."""
        return format_mentors(self.find(discipline_name))


def format_mentors(teachers: list[str]) -> str | None:
    return ", ".join(teachers) if teachers else None


def _find_columns(ws) -> tuple[int, int, int]:
    """
    Определяет индексы столбцов (образовательная программа, дисциплина, ФИО)
    по заголовкам первых строк листа; если заголовки не распознаны -
    используется стандартная раскладка A, B, D.
    """
    program_col, discipline_col, teacher_col = 0, 1, 3
    for row in ws.iter_rows(min_row=1, max_row=5):
        for idx, cell in enumerate(row):
            if not isinstance(cell.value, str):
                continue
            text = normalize_text(cell.value)
            if "специальность" in text:
                program_col = idx
            elif text.startswith("наименование"):
                discipline_col = idx
            elif text.startswith("ф и о"):
                teacher_col = idx
    return program_col, discipline_col, teacher_col


def _clean(value) -> str:
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def get_faculty_index(pps_xlsx: ExcelSource, program_name: str) -> FacultyIndex:
    """
    Читает файл ППС и возвращает преподавателей образовательной программы
    program_name (например "Информационные технологии и искусственный
    интеллект"). Программа ищется по названию тоже нечётко - в файле
    встречаются лишние пробелы и т.п.

    :raises ValueError: если такой программы в файле нет
    """
    wb = openpyxl.load_workbook(pps_xlsx, data_only=True, read_only=True)
    try:
        ws = wb.worksheets[0]
        program_col, discipline_col, teacher_col = _find_columns(ws)

        rows: list[tuple[str, str, str]] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            width = max(program_col, discipline_col, teacher_col) + 1
            if len(row) < width:
                continue
            program = _clean(row[program_col])
            discipline = _clean(row[discipline_col])
            teacher = _clean(row[teacher_col])
            if program and discipline and teacher:
                rows.append((program, discipline, teacher))
    finally:
        wb.close()

    programs = {normalize_text(program) for program, _, _ in rows}
    target = normalize_text(program_name)
    if target in programs:
        matched_program = target
    else:
        match = difflib.get_close_matches(
            target, sorted(programs), n=1, cutoff=PROGRAM_CUTOFF
        )
        if not match:
            raise ValueError(
                f'Образовательная программа "{program_name}" не найдена в файле ППС'
            )
        matched_program = match[0]

    teachers: dict[str, list[str]] = {}
    for program, discipline, teacher in rows:
        if normalize_text(program) != matched_program:
            continue
        names = teachers.setdefault(normalize_text(discipline), [])
        if teacher not in names:
            names.append(teacher)

    return FacultyIndex(teachers)
